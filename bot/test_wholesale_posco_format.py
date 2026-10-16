from __future__ import annotations

import base64
import os
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from posco_order_transform import HEADERS
from wholesale_posco_format import generate_posco_from_po_new


class WholesalePoscoFormatTests(unittest.TestCase):
    def _template_secret(self) -> str:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "New Order"
        sheet.append(list(HEADERS))
        sheet.append([None] * len(HEADERS))
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as temporary:
            path = Path(temporary.name)
        try:
            workbook.save(path)
            return base64.b64encode(path.read_bytes()).decode("ascii")
        finally:
            path.unlink(missing_ok=True)
            workbook.close()

    def _source_workbook(self, path: Path) -> None:
        workbook = Workbook()
        order = workbook.active
        order.title = "PO Nuevo"
        order["A9"] = "5545156"
        order["F11"] = "Mex y Can USA"
        order["F15"] = "TPI"
        order["B19"] = datetime(2026, 10, 15)
        order["B22"] = "WS MARIGOLD ASSORTED 5ST PK10"
        order["C22"] = "ORANGE"
        order["D22"] = 5
        order["E22"] = 10
        order["F22"] = "F4P"
        order["B23"] = "Manual White Product"
        order["C23"] = "WHITE"
        order["D23"] = 6
        order["E23"] = 12
        order["F23"] = "F4P"
        order["J23"] = "PENDIENTE DE BOOM"
        products = workbook.create_sheet("Productos")
        products.append([
            "Distribuidor", "Cust Name", "CATEGORIA", "FLOR", "Flor Color",
            "Descripcion", "Ramos/Cajas", "txr_orden", "Caja",
        ])
        products.append([
            "PACIFICA PRODUCE FARMS", "WHOLESALE", "CB / BULK", "MARIGOLD",
            "ORANGE", "WS MARIGOLD ASSORED 5ST PK10", "5", "10", "F4P",
        ])
        workbook.save(path)
        workbook.close()

    def test_creates_posco_rows_and_flags_only_unmapped_products(self):
        previous = os.environ.get("KOMET_POSCO_TEMPLATE_B64")
        os.environ["KOMET_POSCO_TEMPLATE_B64"] = self._template_secret()
        try:
            with tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "source.xlsx"
                output = Path(directory) / "posco.xlsx"
                self._source_workbook(source)
                result = generate_posco_from_po_new(source, output)
                self.assertEqual(result["rows"], 2)
                self.assertEqual(result["no"], "Mex y Can USA - TPI 5545156")
                self.assertEqual(result["boom_pending"], [{
                    "product": "Manual White Product",
                    "recommendation": "Manual White Product WHITE STEM/6 BU/12 F4P",
                }])
                workbook = load_workbook(output, data_only=True)
                try:
                    sheet = workbook["New Order"]
                    self.assertEqual(sheet["A2"].value, "PACIFICA PRODUCE FARMS")
                    self.assertEqual(sheet["B2"].value, "WHOLESALE")
                    self.assertEqual(sheet["C2"].value, "WS MARIGOLD ASSORED 5ST PK10")
                    self.assertEqual(sheet["F2"].value, "MARIGOLD")
                    self.assertEqual(sheet["G2"].value, "ORANGE")
                    self.assertEqual(sheet["M2"].value, "F4P")
                    self.assertEqual(sheet["F3"].value, "Manual White Product")
                finally:
                    workbook.close()
        finally:
            if previous is None:
                os.environ.pop("KOMET_POSCO_TEMPLATE_B64", None)
            else:
                os.environ["KOMET_POSCO_TEMPLATE_B64"] = previous


if __name__ == "__main__":
    unittest.main()
