from datetime import date
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import types
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

playwright_stub = types.ModuleType("playwright")
playwright_sync_api_stub = sys.modules.setdefault(
    "playwright.sync_api",
    types.ModuleType("playwright.sync_api"),
)
playwright_sync_api_stub.Page = object
playwright_sync_api_stub.TimeoutError = TimeoutError
playwright_sync_api_stub.PlaywrightTimeoutError = TimeoutError
playwright_sync_api_stub.sync_playwright = lambda: None
sys.modules.setdefault("playwright", playwright_stub)
import inventory_boxes
from inventory_box_transform import DATE_NUMBER_FORMAT, normalize_date_formats


class FakeCheckbox:
    def __init__(self, checked: bool) -> None:
        self.checked = checked
        self.clicks = 0

    def is_checked(self) -> bool:
        return self.checked

    def click(self, timeout: int) -> None:
        self.clicks += 1
        self.checked = not self.checked


class FakePage:
    def __init__(self) -> None:
        self.waits: list[int] = []

    def wait_for_timeout(self, milliseconds: int) -> None:
        self.waits.append(milliseconds)


class InventoryBoxesTests(unittest.TestCase):
    @patch("inventory_boxes.awb_checkbox")
    @patch("inventory_boxes.inventory_is_empty", return_value=False)
    @patch("inventory_boxes.selected_inventory_rows", side_effect=[0, 79])
    def test_reselects_rows_when_awb_header_keeps_stale_checked_state(
        self,
        selected_rows,
        inventory_empty,
        find_awb,
    ) -> None:
        checkbox = FakeCheckbox(checked=True)
        find_awb.return_value = checkbox
        page = FakePage()

        inventory_boxes.select_all_inventory(page)

        self.assertEqual(checkbox.clicks, 2)
        self.assertTrue(checkbox.is_checked())

    @patch("inventory_boxes.awb_checkbox", return_value=FakeCheckbox(checked=False))
    @patch("inventory_boxes.inventory_is_empty", return_value=False)
    @patch("inventory_boxes.inventory_processing_visible", side_effect=[True, False])
    def test_waits_for_komet_processing_overlay_to_disappear(
        self,
        processing_visible,
        inventory_empty,
        find_awb,
    ) -> None:
        page = FakePage()

        inventory_boxes.wait_for_inventory_ready(page)

        self.assertEqual(page.waits, [500])

    def test_normalizes_dates_only_on_the_availability_sheet(self) -> None:
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            output = Path(temp_dir) / "normalized.xlsx"
            workbook = Workbook()
            active = workbook.active
            active.title = "Customer View"
            active["A1"] = "Debe conservarse"
            availability = workbook.create_sheet("Availability")
            availability.append(["Product Description", "Available From"])
            availability.append(["A", date(2026, 9, 14)])
            availability["B2"].number_format = "m/d/yy"
            workbook.save(source)
            workbook.close()

            normalize_date_formats(source, output)

            result = load_workbook(output, data_only=False)
            try:
                self.assertEqual(result["Availability"]["B2"].number_format, DATE_NUMBER_FORMAT)
                self.assertEqual(result["Customer View"]["A1"].value, "Debe conservarse")
            finally:
                result.close()


if __name__ == "__main__":
    unittest.main()
