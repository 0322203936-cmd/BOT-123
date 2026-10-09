"""Build review-only POSCO order workbooks from Komet's Order Details XLS."""

from __future__ import annotations

import base64
import binascii
import json
import os
import re
from copy import copy
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path

import xlrd
from openpyxl import load_workbook


HEADERS = (
    "Distribuidor", "Cust Name", "Descripcion", "No.", "CATEGORIA", "FLOR",
    "Flor Color", "Load Date", "#cajas", "Pack", "Stems", "Fecha Produccion", "Caja",
)
BOX_CODES = {"D": "F3", "L": "F4", "1/2L": "F5", "WET": "H3"}
MIXED_BOXES = {"day of dead mixed box", "california greens mixed box"}
CARRIER_ABBREVIATIONS = {"floral trade distributors": "FTD"}


class TransformationError(ValueError):
    pass


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def carrier_abbreviation(value: object) -> str:
    carrier = clean(value)
    if not carrier:
        raise TransformationError("Carrier vacío en el detalle de la orden.")
    known = CARRIER_ABBREVIATIONS.get(carrier.casefold())
    if known:
        return known
    first_word = re.search(r"[A-Za-z]+", carrier)
    if not first_word or len(first_word.group()) < 3:
        raise TransformationError(f"Carrier {carrier!r} no permite obtener una abreviación de tres letras.")
    return first_word.group()[:3].upper()


def required_positive_int(value: object, label: str) -> int:
    if isinstance(value, bool):
        raise TransformationError(f"{label} debe ser un entero positivo; llegó {value!r}.")
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise TransformationError(f"{label} debe ser un entero positivo; llegó {value!r}.") from error
    if not number.is_integer() or number <= 0:
        raise TransformationError(f"{label} debe ser un entero positivo; llegó {value!r}.")
    return int(number)


def optional_positive_int(value: object, label: str) -> int | None:
    return None if clean(value) == "" else required_positive_int(value, label)


def secret_workbook(name: str) -> bytes:
    encoded = os.environ.get(name, "").strip()
    if not encoded:
        raise TransformationError(f"Falta el secreto {name} para generar el artefacto POSCO.")
    try:
        return base64.b64decode(encoded, validate=True)
    except binascii.Error as error:
        raise TransformationError(f"El secreto {name} no contiene un Excel codificado válido.") from error


def load_homologation(workbook_bytes: bytes) -> dict[str, list[dict]]:
    workbook = None
    try:
        workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        headings = [clean(sheet.cell(1, column).value).casefold() for column in range(1, 8)]
        expected = ["flor", "flor color", "producto komet", "stems/bunch", "bunches", "flor posco", "box"]
        if headings != expected:
            raise TransformationError("La homologación no tiene las columnas A:G esperadas.")
        lookup: dict[str, list[dict]] = {}
        for row in sheet.iter_rows(min_row=2, max_col=7, values_only=True):
            flower, color, product, stems, bunches, description, box = row
            if not clean(product):
                continue
            if not all(clean(value) for value in (flower, color, description, box)):
                raise TransformationError(f"Homologación incompleta para {clean(product)}.")
            name = clean(product).casefold()
            mixed = name in MIXED_BOXES
            entry = {
                "flower": clean(flower),
                "color": clean(color),
                "description": clean(description),
                "box": clean(box).upper(),
                "stems": optional_positive_int(stems, "Stems/Bunch de homologación") if not mixed else None,
                "bunches": optional_positive_int(bunches, "Bunches de homologación") if not mixed else None,
            }
            if not mixed and (entry["stems"] is None or entry["bunches"] is None):
                raise TransformationError(f"Faltan Pack o Stems en la homologación de {clean(product)}.")
            lookup.setdefault(name, []).append(entry)
        if not lookup:
            raise TransformationError("La homologación no contiene productos.")
        return lookup
    except TransformationError:
        raise
    except Exception as error:
        raise TransformationError("No se pudo leer el Excel de homologación.") from error
    finally:
        if workbook is not None:
            workbook.close()


def parse_order_xls(path: Path) -> tuple[dict[str, str], list[dict]]:
    try:
        workbook = xlrd.open_workbook(str(path))
        sheet = workbook.sheet_by_name("Order Details") if "Order Details" in workbook.sheet_names() else workbook.sheet_by_index(0)
    except Exception as error:
        raise TransformationError("No se pudo leer el Excel Order Details descargado.") from error
    labels = {
        "order number": "order_number", "customer": "customer",
        "ship date": "ship_date", "carrier": "carrier",
    }
    details: dict[str, object] = {}
    header_row = None
    columns: dict[str, int] = {}
    for row_index in range(sheet.nrows):
        first = clean(sheet.cell_value(row_index, 0))
        second = clean(sheet.cell_value(row_index, 1)) if sheet.ncols > 1 else ""
        field = labels.get(first.rstrip(":").casefold())
        if field:
            details[field] = second
        if first.casefold() == "box code" and second.casefold() == "product description":
            header_row = row_index
            columns = {clean(sheet.cell_value(row_index, column)).casefold(): column for column in range(sheet.ncols)}
            break
    required = {"box code", "product description", "box type", "bunches", "stems/bunch"}
    if header_row is None or not required.issubset(columns):
        raise TransformationError("El XLS no contiene la tabla Boxes con sus columnas esperadas.")
    if not all(details.get(field) for field in labels.values()):
        raise TransformationError("El XLS no contiene Order Number, Customer, Ship Date y Carrier completos.")
    try:
        details["ship_date_value"] = datetime.strptime(details["ship_date"], "%m/%d/%Y")
    except ValueError as error:
        raise TransformationError(f"Ship Date inválido: {details['ship_date']!r}.") from error
    boxes: list[dict] = []
    seen_codes: set[str] = set()
    for row_index in range(header_row + 1, sheet.nrows):
        def field(name: str) -> object:
            return sheet.cell_value(row_index, columns[name])

        code = clean(field("box code"))
        product = clean(field("product description"))
        if not code and not product:
            continue
        if not code or not product:
            raise TransformationError(f"Caja incompleta en la fila {row_index + 1} del XLS.")
        if code in seen_codes:
            raise TransformationError(f"Box Code duplicado en el XLS: {code}.")
        seen_codes.add(code)
        boxes.append({
            "row": row_index + 1,
            "code": code,
            "product": product,
            "box_type": clean(field("box type")).upper(),
            "bunches": field("bunches"),
            "stems": field("stems/bunch"),
        })
    if not boxes:
        raise TransformationError("El XLS Order Details no contiene cajas para transformar.")
    return details, boxes


def build_order_rows(details: dict, boxes: list[dict], lookup: dict[str, list[dict]]) -> list[list]:
    output: list[list] = []
    carrier_code = carrier_abbreviation(details.get("carrier"))
    for box in boxes:
        product = box["product"]
        name = product.casefold()
        box_type = box["box_type"]
        if box_type not in BOX_CODES:
            raise TransformationError(
                f"Orden {details['order_number']}, fila {box['row']}, {product}: "
                f"Box Type {box_type or '(vacío)'} no tiene código de Caja configurado."
            )
        mixed = name in MIXED_BOXES
        bunches = optional_positive_int(box["bunches"], "Bunches") if mixed else required_positive_int(box["bunches"], "Bunches")
        stems = optional_positive_int(box["stems"], "Stems/Bunch") if mixed else required_positive_int(box["stems"], "Stems/Bunch")
        candidates = [entry for entry in lookup.get(name, []) if entry["box"] == box_type]
        if not mixed:
            candidates = [entry for entry in candidates if entry["bunches"] == bunches and entry["stems"] == stems]
        if len(candidates) != 1:
            raise TransformationError(
                f"Orden {details['order_number']}, fila {box['row']}, {product}: "
                f"homologación {'ambigua' if candidates else 'no encontrada'} "
                f"para Box Type={box_type}, Bunches={bunches}, Stems/Bunch={stems}."
            )
        entry = candidates[0]
        ship_date = details["ship_date_value"]
        output.append([
            "PACIFICA PRODUCE FARMS", "WHOLESALE", entry["description"],
            f"{details['customer']} - {carrier_code} {details['order_number']}", "CB / BULK",
            entry["flower"], entry["color"], ship_date, 1, bunches, stems,
            ship_date - timedelta(days=3), BOX_CODES[box_type],
        ])
    return output


def validate_template(template_bytes: bytes) -> None:
    try:
        workbook = load_workbook(BytesIO(template_bytes), read_only=True)
        try:
            sheet = workbook["New Order"]
            if tuple(clean(sheet.cell(1, column).value) for column in range(1, 14)) != HEADERS:
                raise TransformationError("La plantilla New Order no tiene las 13 columnas esperadas.")
        finally:
            workbook.close()
    except TransformationError:
        raise
    except Exception as error:
        raise TransformationError("No se pudo leer la plantilla New Order.") from error


def write_order_workbook(template_bytes: bytes, rows: list[list], destination: Path) -> None:
    try:
        workbook = load_workbook(BytesIO(template_bytes))
        sheet = workbook["New Order"]
        if tuple(clean(sheet.cell(1, column).value) for column in range(1, 14)) != HEADERS:
            raise TransformationError("La plantilla New Order no tiene las 13 columnas esperadas.")
        example_styles = [copy(sheet.cell(2, column)._style) for column in range(1, 14)]
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.value = None
        for row_index, values in enumerate(rows, start=2):
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_index, column)
                cell._style = copy(example_styles[column - 1])
                cell.value = value
                if column in (8, 12):
                    cell.number_format = "mm-dd-yy"
        if sheet.auto_filter.ref:
            sheet.auto_filter.ref = f"A1:M{len(rows) + 1}"
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
        check = load_workbook(destination, read_only=True, data_only=True)
        try:
            check_sheet = check["New Order"]
            for row_index, values in enumerate(rows, start=2):
                if [check_sheet.cell(row_index, column).value for column in range(1, 14)] != values:
                    raise TransformationError("La verificación del Excel generado no coincide con las cajas de origen.")
        finally:
            check.close()
    except TransformationError:
        destination.unlink(missing_ok=True)
        raise
    except Exception as error:
        destination.unlink(missing_ok=True)
        raise TransformationError("No se pudo crear o verificar el Excel POSCO.") from error


def load_order_state(path: Path) -> tuple[set[str], set[str]]:
    if not path.exists():
        return set(), set()
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        transformed = payload["transformed_orders"]
        ignored = payload.get("ignored_orders", [])
        if (
            payload.get("version") != 1
            or not isinstance(transformed, list)
            or not isinstance(ignored, list)
            or any(not isinstance(key, str) for key in transformed + ignored)
            or set(transformed) & set(ignored)
        ):
            raise ValueError("formato inválido")
        return set(transformed), set(ignored)
    except (OSError, ValueError, KeyError, TypeError) as error:
        raise TransformationError("No se pudo leer la bitácora separada de transformaciones.") from error


def load_transformed_keys(path: Path) -> set[str]:
    return load_order_state(path)[0]


def save_transformed_keys(keys: set[str], path: Path) -> None:
    _, ignored = load_order_state(path)
    if keys & ignored:
        raise TransformationError("Una orden manual no puede marcarse como transformada.")
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(
            {"version": 1, "transformed_orders": sorted(keys), "ignored_orders": sorted(ignored)}, indent=2
        ) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)
