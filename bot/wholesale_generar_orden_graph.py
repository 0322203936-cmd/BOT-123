"""Generate the Wholesale order PDF through Microsoft Graph.

The former implementation opened Excel Online and clicked the Office Script
button.  Graph cannot invoke an Office Script directly, so this bot reproduces
the script's observable behavior safely:

* download the source workbook through Graph;
* create a short-lived workbook copy with only ``PO Nuevo`` visible;
* apply the same print area, orientation, and letter paper settings;
* ask Graph to convert that temporary workbook to PDF; and
* delete the temporary SharePoint item in a ``finally`` block.

The permanent XLSM is never overwritten and no interactive SharePoint login is
needed.  The existing workbook, macros, and other bots are left untouched.
"""

from __future__ import annotations

import json
import os
import tempfile
import time
from pathlib import Path
from urllib.parse import quote
from uuid import uuid4

import requests
from openpyxl import load_workbook

from sharepoint_sync import (
    GRAPH_URL,
    graph_headers,
    graph_token,
    resolve_sharepoint_item_by_url,
)


DEFAULT_WORKBOOK_URL = (
    "https://pacificafarms.sharepoint.com/:x:/r/sites/requerimientovsproyeccion/"
    "_layouts/15/Doc.aspx?sourcedoc=%7B7FEB8241-AC2E-4964-B671-C4FAFF942F33%7D&"
    "file=PEDIDO%20Suzy%20-%20Mex%20y%20Can%20USA%2008-18-2026.xlsm&"
    "action=edit&mobileredirect=true"
)
TARGET_SHEET = "PO Nuevo"
PRINT_AREA = "A1:I40"
ARTIFACT_DIR = Path(
    os.environ.get(
        "WHOLESALE_ORDER_ARTIFACT_DIR", "artifacts/wholesale_generar_orden"
    )
)


def request_graph(
    method: str,
    url: str,
    headers: dict[str, str],
    *,
    data: bytes | None = None,
    timeout: int = 180,
) -> requests.Response:
    retryable = {409, 423, 429, 500, 502, 503, 504}
    for attempt in range(1, 6):
        response = requests.request(
            method,
            url,
            headers=headers,
            data=data,
            timeout=timeout,
            allow_redirects=True,
        )
        if response.status_code in retryable and attempt < 5:
            retry_after = int(response.headers.get("Retry-After", "0") or 0)
            time.sleep(retry_after or min(attempt * 5, 30))
            continue
        if not response.ok:
            try:
                detail = response.json().get("error", {}).get("message", "")
            except ValueError:
                detail = response.text[:500]
            raise RuntimeError(f"Microsoft Graph respondió {response.status_code}: {detail}")
        return response
    raise RuntimeError("Microsoft Graph no respondió después de los reintentos.")


def download_workbook(token: str, item: dict, destination: Path) -> None:
    drive_id = item["parentReference"]["driveId"]
    response = request_graph(
        "GET",
        f"{GRAPH_URL}/drives/{drive_id}/items/{item['id']}/content",
        graph_headers(token),
        timeout=180,
    )
    destination.write_bytes(response.content)


def prepare_pdf_copy(source: Path, destination: Path) -> dict[str, object]:
    """Mirror the Office Script's workbook-only PDF preparation."""
    workbook = load_workbook(source, keep_vba=True, keep_links=True, data_only=False)
    try:
        if TARGET_SHEET not in workbook.sheetnames:
            raise RuntimeError(f"El libro no contiene la hoja {TARGET_SHEET!r}.")

        target_index = workbook.sheetnames.index(TARGET_SHEET)
        target = workbook[TARGET_SHEET]

        # Equivalent to setVisibility(visible), activate(), and hiding every
        # other sheet in the Office Script.
        for sheet in workbook.worksheets:
            sheet.sheet_state = "visible" if sheet is target else "hidden"
            sheet.sheet_view.tabSelected = sheet is target

        # Equivalent to layout.setPrintArea("A1:I40"), portrait, letter.
        target.print_area = PRINT_AREA
        target.page_setup.orientation = target.ORIENTATION_PORTRAIT
        target.page_setup.paperSize = target.PAPERSIZE_LETTER
        workbook.views[0].activeTab = target_index

        workbook.save(destination)
        return {
            "sheet": TARGET_SHEET,
            "print_area": PRINT_AREA,
            "orientation": "portrait",
            "paper": "letter",
            "sheet_count": len(workbook.sheetnames),
        }
    finally:
        workbook.close()


def upload_temporary_workbook(token: str, item: dict, workbook_path: Path) -> dict:
    drive_id = item["parentReference"]["driveId"]
    parent_id = item["parentReference"]["id"]
    filename = f".codex-generar-orden-{uuid4().hex}.xlsm"
    upload_url = (
        f"{GRAPH_URL}/drives/{drive_id}/items/{parent_id}:/{quote(filename, safe='')}:/content"
    )
    response = request_graph(
        "PUT",
        upload_url,
        {**graph_headers(token), "Content-Type": "application/octet-stream"},
        data=workbook_path.read_bytes(),
        timeout=180,
    )
    temporary_item = response.json()
    temporary_item["_temporary_name"] = filename
    return temporary_item


def convert_to_pdf(token: str, item: dict, destination: Path) -> None:
    drive_id = item["parentReference"]["driveId"]
    response = request_graph(
        "GET",
        f"{GRAPH_URL}/drives/{drive_id}/items/{item['id']}/content?format=pdf",
        graph_headers(token),
        timeout=240,
    )
    if not response.content.startswith(b"%PDF"):
        raise RuntimeError("Graph no devolvió un PDF válido al convertir el libro.")
    destination.write_bytes(response.content)


def delete_temporary_workbook(token: str, item: dict) -> None:
    drive_id = item["parentReference"]["driveId"]
    request_graph(
        "DELETE",
        f"{GRAPH_URL}/drives/{drive_id}/items/{item['id']}",
        graph_headers(token),
        timeout=60,
    )


def main() -> None:
    ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    started = time.time()
    summary: dict[str, object] = {
        "button": "Ejecutar Generar Orden",
        "engine": "Microsoft Graph",
        "workbook": os.environ.get("SHAREPOINT_WORKBOOK_URL", DEFAULT_WORKBOOK_URL),
        "status": "error",
    }
    temporary_item: dict | None = None

    try:
        token = graph_token()
        item = resolve_sharepoint_item_by_url(
            token, os.environ.get("SHAREPOINT_WORKBOOK_URL", DEFAULT_WORKBOOK_URL)
        )
        with tempfile.TemporaryDirectory(prefix="wholesale-generar-orden-") as temp_dir:
            source = Path(temp_dir) / "source.xlsm"
            prepared = Path(temp_dir) / "prepared.xlsm"
            download_workbook(token, item, source)
            summary["preparation"] = prepare_pdf_copy(source, prepared)

            temporary_item = upload_temporary_workbook(token, item, prepared)
            summary["temporary_sharepoint_item"] = temporary_item.get("name")
            pdf_path = ARTIFACT_DIR / "generar-orden.pdf"
            convert_to_pdf(token, temporary_item, pdf_path)
            summary.update(
                {
                    "status": "completed",
                    "pdf": pdf_path.name,
                    "pdf_bytes": pdf_path.stat().st_size,
                    "elapsed_seconds": round(time.time() - started, 1),
                }
            )
            print(f"PDF generado con Graph: {pdf_path}", flush=True)
    except Exception as error:
        summary.update(
            {"error": str(error), "elapsed_seconds": round(time.time() - started, 1)}
        )
        raise
    finally:
        if temporary_item:
            try:
                delete_temporary_workbook(token, temporary_item)
                summary["temporary_sharepoint_item_deleted"] = True
            except Exception as cleanup_error:
                summary["temporary_sharepoint_item_deleted"] = False
                summary["cleanup_error"] = str(cleanup_error)
        (ARTIFACT_DIR / "summary.json").write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
        )


if __name__ == "__main__":
    main()
