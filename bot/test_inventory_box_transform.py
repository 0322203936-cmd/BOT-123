from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table
from openpyxl.styles import Font, PatternFill

try:
    from inventory_box_transform import (
        apply_inventory_rules,
        create_single_sheet_workbook,
        transform_inventory_workbook,
        refresh_workbook_with_komet_inventory,
    )
except ImportError:
    try:
        from bot.inventory_box_transform import (
            apply_inventory_rules,
            create_single_sheet_workbook,
            transform_inventory_workbook,
            refresh_workbook_with_komet_inventory,
        )
    except ImportError:
        apply_inventory_rules = None
        create_single_sheet_workbook = None
        transform_inventory_workbook = None
        refresh_workbook_with_komet_inventory = None


class InventoryBoxTransformTests(unittest.TestCase):
    assumed_today = date(2026, 9, 10)

    def test_removes_window_and_sundays_and_adds_one_row_per_product(self) -> None:
        rows = [
            ["Pacific", "A", "Bunch", date(2026, 9, 9)],
            ["Pacific", "A", "Bunch", date(2026, 9, 10)],
            ["Pacific", "A", "Bunch", date(2026, 9, 14)],
            ["Pacific", "B", "Bunch", date(2026, 9, 13)],
            ["Pacific", "B", "Bunch", date(2026, 9, 15)],
        ]

        self.assertIsNotNone(apply_inventory_rules)
        result = apply_inventory_rules(
            rows,
            product_column=1,
            date_column=3,
            assumed_today=self.assumed_today,
        )

        output_dates = [row.values[3] for row in result.output_rows]
        self.assertEqual(
            output_dates,
            [
                date(2026, 9, 9),
                date(2026, 9, 14),
                date(2026, 9, 15),
                date(2026, 9, 15),
                date(2026, 9, 16),
            ],
        )
        added = [row for row in result.output_rows if row.is_added]
        self.assertEqual(len(added), 2)
        self.assertEqual([row.values[1] for row in added], ["A", "B"])
        self.assertEqual(result.removed_rows, 2)
        self.assertEqual(result.final_sunday_rows, 0)

    def test_saturday_latest_row_adds_monday(self) -> None:
        rows = [["Pacific", "A", "Bunch", date(2026, 9, 19)]]

        self.assertIsNotNone(apply_inventory_rules)
        result = apply_inventory_rules(
            rows,
            product_column=1,
            date_column=3,
            assumed_today=self.assumed_today,
        )

        self.assertEqual(result.output_rows[-1].values[3], date(2026, 9, 21))

    def test_rejects_a_product_with_no_surviving_dated_row(self) -> None:
        rows = [["Pacific", "A", "Bunch", date(2026, 9, 10)]]

        self.assertIsNotNone(apply_inventory_rules)
        with self.assertRaisesRegex(RuntimeError, "A"):
            apply_inventory_rules(
                rows,
                product_column=1,
                date_column=3,
                assumed_today=self.assumed_today,
            )

    def test_workbook_transform_copies_format_to_added_row(self) -> None:
        self.assertIsNotNone(transform_inventory_workbook)
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            output = Path(temp_dir) / "output.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Availability"
            sheet.append(["Vendor Name", "Product Description", "Package Type", "Available From"])
            sheet.append(["Pacific", "A", "Bunch", date(2026, 9, 14)])
            sheet["B2"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
            sheet["B2"].font = Font(name="Arial", bold=True)
            sheet["D2"].number_format = r"yyyy\-mm\-dd"
            workbook.save(source)
            workbook.close()

            transform_inventory_workbook(
                source,
                output,
                assumed_today=self.assumed_today,
            )

            result = load_workbook(output, data_only=False)
            try:
                sheet = result.active
                self.assertEqual(sheet.max_row, 3)
                self.assertEqual(sheet["B3"].value, "A")
                self.assertEqual(sheet["D3"].value.date(), date(2026, 9, 15))
                self.assertEqual(sheet["B3"]._style, sheet["B2"]._style)
                self.assertEqual(sheet["D3"].number_format, sheet["D2"].number_format)
            finally:
                result.close()

    def test_creates_a_komet_copy_with_only_the_availability_sheet(self) -> None:
        self.assertIsNotNone(create_single_sheet_workbook)
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            output = Path(temp_dir) / "komet.xlsx"
            workbook = Workbook()
            active = workbook.active
            active.title = "Customer View"
            active["A1"] = "No debe ir a Komet"
            availability = workbook.create_sheet("Availability")
            availability.append(["Product Description", "Available From"])
            availability.append(["A", date(2026, 9, 14)])
            availability["A2"].font = Font(name="Arial", bold=True)
            workbook.save(source)
            workbook.close()

            create_single_sheet_workbook(source, output)

            result = load_workbook(output, data_only=False)
            try:
                self.assertEqual(result.sheetnames, ["Availability"])
                self.assertEqual(result.active["A2"].value, "A")
                self.assertEqual(result.active["A2"]._style, availability["A2"]._style)
            finally:
                result.close()

    def test_keeps_the_second_sheet_in_the_full_transformed_workbook(self) -> None:
        self.assertIsNotNone(transform_inventory_workbook)
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            output = Path(temp_dir) / "output.xlsx"
            workbook = Workbook()
            active = workbook.active
            active.title = "Customer View"
            active["A1"] = "Debe conservarse"
            availability = workbook.create_sheet("Availability")
            availability.append(["Product Description", "Available From"])
            availability.append(["A", date(2026, 9, 14)])
            workbook.save(source)
            workbook.close()

            transform_inventory_workbook(source, output, assumed_today=self.assumed_today)

            result = load_workbook(output, data_only=False)
            try:
                self.assertEqual(result.sheetnames, ["Customer View", "Availability"])
                self.assertEqual(result["Customer View"]["A1"].value, "Debe conservarse")
                self.assertEqual(result["Availability"]["A3"].value, "A")
            finally:
                result.close()

    def test_verifies_added_rows_against_their_retained_source_rows(self) -> None:
        self.assertIsNotNone(transform_inventory_workbook)
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            output = Path(temp_dir) / "output.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Availability"
            sheet.append(["Product Description", "Available From"])
            sheet.append(["A", date(2026, 9, 14)])
            sheet.append(["B", date(2026, 9, 10)])
            sheet.append(["B", date(2026, 9, 15)])
            sheet["A2"].fill = PatternFill(fill_type="solid", fgColor="FF0000")
            sheet["A4"].fill = PatternFill(fill_type="solid", fgColor="00FF00")
            workbook.save(source)
            workbook.close()

            transform_inventory_workbook(source, output, assumed_today=self.assumed_today)

            result = load_workbook(output, data_only=False)
            try:
                self.assertEqual(result.active["A4"].value, "A")
                self.assertEqual(result.active["A5"].value, "B")
                self.assertEqual(result.active["A5"]._style, result.active["A3"]._style)
            finally:
                result.close()

    def test_replaces_inventory_and_updates_availability_without_increasing_qty(self) -> None:
        self.assertIsNotNone(refresh_workbook_with_komet_inventory)
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            komet = Path(temp_dir) / "komet.xlsx"
            output = Path(temp_dir) / "output.xlsx"

            workbook = Workbook()
            customer = workbook.active
            customer.title = "Customer View"
            customer["A1"] = "=SUM(tblAvailability[Qty Packages])"
            availability = workbook.create_sheet("Availability")
            availability.append(
                [
                    "Vendor Name",
                    "Product Description",
                    "Unit of Sale",
                    "Package Type",
                    "Pack",
                    "Units / Pack",
                    "Qty Packages",
                    "Price",
                    "Available From",
                ]
            )
            availability.append(["Pacific", "A product", "Bunch", "L", 10, 1, "=OLD", 3, date(2026, 9, 15)])
            availability.append(["Pacific", "B product", "Bunch", "L", 10, 1, 3, 3, date(2026, 9, 15)])
            availability.append(["Pacific", "Missing product", "Bunch", "L", 10, 1, "=OLD", 3, date(2026, 9, 15)])
            availability.add_table(Table(displayName="tblAvailability", ref="A1:I4"))

            inventory = workbook.create_sheet("Inventory")
            inventory.append(["Inventory source"])
            for _ in range(6):
                inventory.append([])
            inventory_headers = ["AWB", "Ref #", "Location", "Product", "Hold", "Customer", "Vendor", "Aging", "Qty"]
            inventory.append(inventory_headers)
            inventory.append(["AWB-2026-09-14", "old", "L", "A product", "", "", "", 1, 8])
            inventory.append(["AWB-2026-09-14", "old", "L", "Missing product", "", "", "", 1, 4])
            inventory.add_table(Table(displayName="tblInventory", ref="A8:I10"))
            workbook.save(source)
            workbook.close()

            export = Workbook()
            export_sheet = export.active
            export_sheet.append(inventory_headers)
            export_sheet.append(["AWB-2026-09-10", "new", "L", "A product", "", "", "", 5, 5])
            export_sheet.append(["AWB-2026-09-10", "new", "L", "B product", "", "", "", 5, 9])
            export.save(komet)
            export.close()

            result = refresh_workbook_with_komet_inventory(
                source,
                komet,
                output,
                assumed_today=self.assumed_today,
            )

            workbook = load_workbook(output, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["Customer View", "Availability", "Inventory"])
                self.assertEqual(workbook["Customer View"]["A1"].value, "=SUM(tblAvailability[Qty Packages])")
                self.assertEqual([workbook["Availability"][f"G{row}"].value for row in (2, 3, 4)], [5, 3, 0])
                self.assertTrue(all(not str(workbook["Availability"][f"G{row}"].value).startswith("=") for row in (2, 3, 4)))
                self.assertEqual(workbook["Inventory"]["D9"].value, "A product")
                self.assertEqual(workbook["Inventory"]["I9"].value, 5)
                self.assertEqual(workbook["Inventory"].tables["tblInventory"].ref, "A8:I10")
            finally:
                workbook.close()

            self.assertEqual(result.updated_rows, 2)
            self.assertEqual(result.decreased_rows, 2)
            self.assertEqual(result.before_total, 15)
            self.assertEqual(result.after_total, 8)


if __name__ == "__main__":
    unittest.main()
