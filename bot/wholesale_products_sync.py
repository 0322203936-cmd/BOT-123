"""Append new WHOLESALE products from the POSCO export to SharePoint.

The export bot remains responsible for downloading the report.  This module
only synchronizes the Productos sheet and is deliberately append-only: existing
descriptions are never rewritten or deleted.
"""

from __future__ import annotations

import math
import os
import re
import time
from pathlib import Path
from typing import Any
from urllib.parse import quote

from openpyxl import load_workbook

from sharepoint_sync import GRAPH_URL, graph_headers, graph_token, resolve_sharepoint_item_by_url


PRODUCT_SHEET = "Productos"
PRODUCT_COLUMN_COUNT = 12  # A:L
WHOLESALE = "WHOLESALE"
FIXED_DISTRIBUTOR = "PACIFICA PRODUCE FARMS"
SHAREPOINT_PRODUCTS_URL = (
    "https://pacificafarms.sharepoint.com/:x:/r/sites/"
    "requerimientovsproyeccion/_layouts/15/Doc.aspx?"
    "sourcedoc=%7BBE147792-3CD4-49FB-8185-63CC0D0833CC%7D&"
    "file=PEDIDO%20Suzy%20-%20Mex%20y%20Can%20USA%2008-18-2026.xlsm&"
    "action=default&mobileredirect=true&wdwpf=doclib-c"
)
RETRYABLE_STATUS = {409, 423, 429, 500, 502, 503, 504}


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalized(value: Any) -> str:
    return clean(value).casefold()


def safe_cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return ""
    return value


def column_letter(number: int) -> str:
    result = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        result = chr(ord("A") + remainder) + result
    return result


def parse_start_cell(address: str) -> tuple[int, int]:
    match = re.search(r"(?:!|^)\$?([A-Z]+)\$?(\d+)", address.upper())
    if not match:
        raise RuntimeError(f"Rango de Productos inválido: {address!r}.")
    column = 0
    for character in match.group(1):
        column = column * 26 + ord(character) - ord("A") + 1
    return column, int(match.group(2))


def header_aliases() -> dict[str, tuple[str, ...]]:
    return {
        "distribuidor": ("distribuidor",),
        "cust_name": ("cust name", "customer", "customer name", "custname"),
        "descripcion": ("descripcion", "description"),
        "categoria": ("categoria", "category"),
        "flor": ("flor", "flower"),
        "flor_color": ("flor color", "flower color", "color"),
        "ramos_cajas": (
            "ramos/cajas",
            "ramos por caja",
            "ramos por cajas",
            "bunches/boxes",
        ),
        "txr_orden": ("txr_orden", "txr orden", "txr"),
        "caja": ("caja", "box", "box type"),
    }


def find_header_row(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    aliases = header_aliases()
    alias_to_key = {
        normalized(alias): key
        for key, values in aliases.items()
        for alias in values
    }
    for row_index, row in enumerate(rows[:15]):
        found: dict[str, int] = {}
        for column_index, value in enumerate(row):
            key = alias_to_key.get(normalized(value))
            if key and key not in found:
                found[key] = column_index
        if "cust_name" in found and "descripcion" in found:
            return row_index, found
    raise RuntimeError(
        "El reporte de POSCO no contiene las columnas Cust Name y Descripcion."
    )


def report_products(report_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(report_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    if not rows:
        return []
    header_index, columns = find_header_row(rows)
    products: list[dict[str, Any]] = []
    seen: set[str] = set()
    for row in rows[header_index + 1 :]:
        def value(key: str) -> Any:
            index = columns.get(key)
            return safe_cell(row[index]) if index is not None and index < len(row) else ""

        if normalized(value("cust_name")) != normalized(WHOLESALE):
            continue
        description = clean(value("descripcion"))
        key = normalized(description)
        if not key or key in seen:
            continue
        seen.add(key)
        products.append(
            {
                "distribuidor": FIXED_DISTRIBUTOR,
                "cust_name": WHOLESALE,
                "categoria": value("categoria"),
                "flor": value("flor"),
                "flor_color": value("flor_color"),
                "descripcion": description,
                "ramos_cajas": value("ramos_cajas"),
                "txr_orden": value("txr_orden"),
                "caja": value("caja"),
            }
        )
    return products


def make_product_row(product: dict[str, Any], dimensions: dict[str, Any]) -> list[Any]:
    """Return the A:L row used by Productos.

    I is Caja.  K (BOX1) mirrors that same source box code so the existing
    PO Nuevo lookup can find its dimension.  L is filled only when the code
    already has a nonblank dimension in Productos; otherwise it remains blank
    for later homologation.
    """
    box = safe_cell(product.get("caja"))
    box_key = normalized(box)
    dimension = dimensions.get(box_key, "") if box_key else ""
    return [
        FIXED_DISTRIBUTOR,
        WHOLESALE,
        safe_cell(product.get("categoria")),
        safe_cell(product.get("flor")),
        safe_cell(product.get("flor_color")),
        safe_cell(product.get("descripcion")),
        safe_cell(product.get("ramos_cajas")),
        safe_cell(product.get("txr_orden")),
        box,
        "",
        box,
        dimension,
    ]


def deduplicate_against_products(
    products: list[dict[str, Any]], existing_descriptions: set[str], dimensions: dict[str, Any]
) -> list[list[Any]]:
    rows: list[list[Any]] = []
    seen = set(existing_descriptions)
    for product in products:
        key = normalized(product.get("descripcion"))
        if not key or key in seen:
            continue
        seen.add(key)
        rows.append(make_product_row(product, dimensions))
    return rows


class ExcelSession:
    def __init__(self, workbook_url: str, headers: dict[str, str]):
        self.workbook_url = workbook_url
        self.headers = headers
        self.session_headers: dict[str, str] | None = None

    def open(self) -> None:
        response = self.request_raw(
            "POST",
            f"{self.workbook_url}/createSession",
            json={"persistChanges": True},
            timeout=120,
            use_session=False,
        )
        session_id = response.json().get("id")
        if not session_id:
            raise RuntimeError("Excel Online no devolvió una sesión válida.")
        self.session_headers = {**self.headers, "workbook-session-id": session_id}

    def close(self) -> None:
        if not self.session_headers:
            return
        try:
            self.request_raw(
                "POST",
                f"{self.workbook_url}/closeSession",
                timeout=30,
            )
        except Exception:
            pass
        finally:
            self.session_headers = None

    def request_raw(
        self,
        method: str,
        url: str,
        *,
        json: dict[str, Any] | None = None,
        timeout: int = 120,
        use_session: bool = True,
    ):
        import requests

        headers = self.session_headers if use_session and self.session_headers else self.headers
        for attempt in range(1, 6):
            response = requests.request(method, url, headers=headers, json=json, timeout=timeout)
            if response.status_code in RETRYABLE_STATUS and attempt < 5:
                retry_after = int(response.headers.get("Retry-After", "0") or 0)
                delay = retry_after or min(attempt * 5, 30)
                print(
                    f"Excel Online respondió HTTP {response.status_code}; "
                    f"reintento {attempt}/4 en {delay} segundos...",
                    flush=True,
                )
                time.sleep(delay)
                continue
            if not response.ok:
                try:
                    detail = response.json().get("error", {}).get("message", "")
                except ValueError:
                    detail = response.text
                raise RuntimeError(f"Microsoft Graph respondió {response.status_code}: {detail}")
            return response
        raise RuntimeError("Excel Online no respondió después de los reintentos.")

    def request(self, method: str, url: str, **kwargs):
        return self.request_raw(method, url, **kwargs)


def worksheet_url(workbook_url: str) -> str:
    return f"{workbook_url}/worksheets('{quote(PRODUCT_SHEET, safe='')}')"


def read_products_state(session: ExcelSession, workbook_url: str) -> tuple[set[str], dict[str, Any], int]:
    used = session.request("GET", f"{worksheet_url(workbook_url)}/usedRange").json()
    address = str(used.get("address") or "")
    values = used.get("values")
    if not address or not isinstance(values, list):
        raise RuntimeError("No se pudo leer el rango usado de Productos.")
    start_column, start_row = parse_start_cell(address)

    header_index = None
    header_columns: dict[str, int] = {}
    alias_to_key = {
        normalized(alias): key
        for key, aliases in header_aliases().items()
        for alias in aliases
    }
    for row_index, row in enumerate(values[:15]):
        mapped = {}
        for offset, value in enumerate(row):
            key = alias_to_key.get(normalized(value))
            if key and key not in mapped:
                mapped[key] = start_column + offset
        if "descripcion" in mapped:
            header_index = row_index
            header_columns = mapped
            break
    if header_index is None:
        raise RuntimeError("No se encontró el encabezado Descripcion en Productos.")

    description_column = header_columns["descripcion"] - start_column
    box_column = header_columns.get("caja", start_column + 8) - start_column
    box1_column = 11 - start_column
    dimension_column = 12 - start_column
    existing: set[str] = set()
    dimensions: dict[str, Any] = {}
    last_data_row = start_row + header_index
    for offset, row in enumerate(values[header_index + 1 :], start=header_index + 1):
        description = row[description_column] if description_column < len(row) else ""
        if clean(description):
            existing.add(normalized(description))
            last_data_row = start_row + offset
        box_value = row[box1_column] if box1_column < len(row) else ""
        dimension = row[dimension_column] if dimension_column < len(row) else ""
        if clean(box_value) and clean(dimension):
            dimensions.setdefault(normalized(box_value), dimension)
        elif box_column < len(row) and clean(row[box_column]) and clean(dimension):
            dimensions.setdefault(normalized(row[box_column]), dimension)
    return existing, dimensions, last_data_row


def find_products_table(session: ExcelSession, workbook_url: str) -> tuple[str, int] | None:
    response = session.request("GET", f"{worksheet_url(workbook_url)}/tables").json()
    for table in response.get("value", []):
        table_id = str(table.get("id") or table.get("name") or "")
        if not table_id:
            continue
        encoded = quote(table_id, safe="")
        try:
            header = session.request(
                "GET",
                f"{worksheet_url(workbook_url)}/tables('{encoded}')/headerRowRange",
            ).json()
        except RuntimeError:
            continue
        values = header.get("values") or []
        if not values:
            continue
        headers = values[0]
        if any(normalized(value) == "descripcion" for value in headers):
            return table_id, len(headers)
    return None


def append_rows(
    session: ExcelSession,
    workbook_url: str,
    rows: list[list[Any]],
    start_row: int,
) -> str:
    table = find_products_table(session, workbook_url)
    if table and table[1] == PRODUCT_COLUMN_COUNT:
        table_id, _ = table
        encoded = quote(table_id, safe="")
        endpoint = f"{worksheet_url(workbook_url)}/tables('{encoded}')/rows/add"
        for offset in range(0, len(rows), 100):
            chunk = rows[offset : offset + 100]
            session.request("POST", endpoint, json={"index": None, "values": chunk})
        return f"tabla {table_id}"

    for offset in range(0, len(rows), 100):
        chunk = rows[offset : offset + 100]
        first = start_row + offset
        last = first + len(chunk) - 1
        address = f"A{first}:L{last}"
        session.request(
            "PATCH",
            f"{worksheet_url(workbook_url)}/range(address='{address}')",
            json={"values": chunk},
        )
    return f"rango A{start_row}:L{start_row + len(rows) - 1}"


def sync_wholesale_products(report_path: Path) -> int:
    products = report_products(report_path)
    token = graph_token()
    workbook_url = os.environ.get("SHAREPOINT_PRODUCTS_URL", SHAREPOINT_PRODUCTS_URL).strip()
    item = resolve_sharepoint_item_by_url(token, workbook_url)
    workbook_api_url = (
        f"{GRAPH_URL}/drives/{item['parentReference']['driveId']}"
        f"/items/{item['id']}/workbook"
    )
    session = ExcelSession(workbook_api_url, {**graph_headers(token), "Content-Type": "application/json"})
    session.open()
    try:
        existing, dimensions, last_data_row = read_products_state(session, workbook_api_url)
        rows = deduplicate_against_products(products, existing, dimensions)
        if not rows:
            print(
                f"Productos WHOLESALE: 0 nuevos; {len(products)} producto(s) del reporte "
                "ya estaban registrados.",
                flush=True,
            )
            return 0
        destination = append_rows(session, workbook_api_url, rows, last_data_row + 1)
        print(
            f"Productos WHOLESALE actualizados: {len(rows)} nuevo(s), "
            f"destino={destination}.",
            flush=True,
        )
        return len(rows)
    finally:
        session.close()


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        raise SystemExit("Uso: python bot/wholesale_products_sync.py <reporte.xlsx>")
    sync_wholesale_products(Path(sys.argv[1]))
