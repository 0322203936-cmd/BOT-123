import io
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from bot.posco_order_transform import (
    HEADERS,
    TransformationError,
    build_order_rows,
    load_homologation,
    load_order_state,
    load_transformed_keys,
    save_transformed_keys,
    validate_template,
    write_order_workbook,
)
from bot import posco_order_artifacts as stage


def workbook_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    if rows and rows[0] == list(HEADERS):
        sheet.title = "New Order"
    for row in rows:
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class PoscoTransformTests(unittest.TestCase):
    def setUp(self):
        self.homologation = load_homologation(workbook_bytes([
            ["FLOR", "FLOR COLOR", "PRODUCTO KOMET", "Stems/Bunch", "Bunches", "FLOR POSCO", "BOX"],
            ["SUNFLOWER TT", "SUNFLOWER TT - PLUM PINK", "Sunflower PLUM TT Pink 60-65 cm", 5, 10,
             "WS SUNFLOWER TT PINK 5 ST PK 10", "L"],
            ["MIXED", "MIXED - GREEN", "California Greens Mixed Box", None, None,
             "WS California Greens Mixed Box PK 10", "L"],
        ]))
        self.template = workbook_bytes([
            list(HEADERS),
            ["SAMPLE VENDOR", "SAMPLE CUSTOMER", "SAMPLE DESCRIPTION", "SAMPLE ORDER"] + [None] * 9,
        ])
        self.details = {
            "order_number": "000317", "customer": "ARIZONA FLORAL EXCHANGE",
            "carrier": "Floral Trade Distributors",
            "ship_date_value": datetime(2026, 10, 12),
        }
        self.box = {
            "row": 11, "code": "568358270", "product": "Sunflower PLUM TT Pink 60-65 cm",
            "box_type": "L", "bunches": 10.0, "stems": 5.0,
        }

    def test_normal_box_creates_exact_template_columns_and_typed_dates(self):
        validate_template(self.template)
        rows = build_order_rows(self.details, [self.box], self.homologation)
        self.assertEqual(rows[0], [
            "PACIFICA PRODUCE FARMS", "WHOLESALE", "WS SUNFLOWER TT PINK 5 ST PK 10",
            "ARIZONA FLORAL EXCHANGE - FTD 000317", "CB / BULK", "SUNFLOWER TT",
            "SUNFLOWER TT - PLUM PINK", datetime(2026, 10, 12), 1, 10, 5,
            datetime(2026, 10, 9), "F4",
        ])
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "order.xlsx"
            write_order_workbook(self.template, rows, destination)
            book = load_workbook(destination, read_only=True)
            sheet = book.active
            self.assertEqual(sheet.title, "New Order")
            self.assertEqual([sheet.cell(1, c).value for c in range(1, 14)], list(HEADERS))
            self.assertEqual(sheet["C2"].value, rows[0][2])
            self.assertEqual(sheet["D2"].value, "ARIZONA FLORAL EXCHANGE - FTD 000317")
            self.assertEqual(sheet["L2"].value, datetime(2026, 10, 9))
            self.assertEqual(sheet["I2"].value, 1)
            self.assertNotEqual(sheet["A2"].value, "SAMPLE VENDOR")
            book.close()

    def test_template_with_wrong_headers_is_rejected_before_order_processing(self):
        book = load_workbook(io.BytesIO(self.template))
        book["New Order"]["A1"] = "Different header"
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        with self.assertRaisesRegex(TransformationError, "13 columnas"):
            validate_template(stream.getvalue())

    def test_each_box_stays_one_row_even_when_product_repeats(self):
        other = {**self.box, "row": 12, "code": "568358271"}
        rows = build_order_rows(self.details, [self.box, other], self.homologation)
        self.assertEqual(len(rows), 2)
        self.assertEqual([row[8] for row in rows], [1, 1])
        self.assertEqual([row[3] for row in rows], ["ARIZONA FLORAL EXCHANGE - FTD 000317"] * 2)

    def test_other_carriers_use_first_three_letters(self):
        armellini = build_order_rows({**self.details, "carrier": "Armellini - Regular"}, [self.box], self.homologation)
        prime = build_order_rows({**self.details, "carrier": "Prime Floral"}, [self.box], self.homologation)
        self.assertEqual(armellini[0][3], "ARIZONA FLORAL EXCHANGE - ARM 000317")
        self.assertEqual(prime[0][3], "ARIZONA FLORAL EXCHANGE - PRI 000317")

    def test_missing_carrier_is_not_silently_omitted(self):
        with self.assertRaisesRegex(TransformationError, "Carrier vacío"):
            build_order_rows({**self.details, "carrier": ""}, [self.box], self.homologation)

    def test_pack_or_stems_mismatch_is_error(self):
        with self.assertRaisesRegex(TransformationError, "homologación no encontrada"):
            build_order_rows(self.details, [{**self.box, "stems": 6}], self.homologation)

    def test_mixed_box_copies_source_pack_and_stems_without_comparison(self):
        mixed = {**self.box, "product": "California Greens Mixed Box", "bunches": 7, "stems": 3}
        row = build_order_rows(self.details, [mixed], self.homologation)[0]
        self.assertEqual(row[2], "WS California Greens Mixed Box PK 10")
        self.assertEqual((row[9], row[10]), (7, 3))

    def test_unconfigured_box_type_is_error(self):
        with self.assertRaisesRegex(TransformationError, "no tiene código de Caja"):
            build_order_rows(self.details, [{**self.box, "box_type": "HAM"}], self.homologation)

    def test_unknown_product_is_error(self):
        with self.assertRaisesRegex(TransformationError, "homologación no encontrada"):
            build_order_rows(self.details, [{**self.box, "product": "Unknown Flower"}], self.homologation)

    def test_separate_ledger_round_trip(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "transformed.json"
            self.assertEqual(load_transformed_keys(path), set())
            save_transformed_keys({"W000317|10/12/2026"}, path)
            self.assertEqual(load_transformed_keys(path), {"W000317|10/12/2026"})

    def test_manual_order_stays_ignored_when_successful_orders_are_saved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "transformed.json"
            path.write_text(
                '{"version": 1, "transformed_orders": [], "ignored_orders": ["W000308|10/09/2026"]}',
                encoding="utf-8",
            )
            save_transformed_keys({"W000317|10/12/2026"}, path)
            self.assertEqual(load_order_state(path), (
                {"W000317|10/12/2026"}, {"W000308|10/09/2026"},
            ))

    def test_manual_order_is_reported_without_download(self):
        orders = [{"order": "W000308", "date": "10/09/2026", "internal_id": "308"}]
        with (
            patch.object(stage, "collect_filtered_orders", return_value=orders),
            patch.object(stage, "download_with_recovery") as download,
        ):
            report = {"orders": []}
            stage.transform_visible_orders(
                None, set(), {"W000308|10/09/2026"}, self.homologation,
                self.template, report, date(2026, 10, 9), date(2026, 10, 19),
            )
        download.assert_not_called()
        self.assertEqual(report["orders"], [{
            "order": "W000308", "key": "W000308|10/09/2026", "status": "manual",
        }])

    def test_wrong_downloaded_order_is_rejected(self):
        with self.assertRaisesRegex(TransformationError, "no de"):
            stage.verify_order_number("W000317", "000316")

    def test_download_reopens_filtered_list_and_retries_after_timeout(self):
        order = {"order": "W000317", "date": "10/12/2026", "internal_id": "317"}
        page = Mock()
        with (
            patch.object(stage, "find_filtered_order", side_effect=[order, order]) as find,
            patch.object(stage, "restore_filtered_orders") as restore,
            patch.object(stage, "capture_download_error") as capture,
            patch.object(stage.komet, "download_export_details", side_effect=[RuntimeError("download timeout"), Path("order.xls")]) as download,
        ):
            result = stage.download_with_recovery(page, order, 1, date(2026, 10, 9), date(2026, 10, 19))
        self.assertEqual(result, Path("order.xls"))
        self.assertEqual(find.call_count, 2)
        self.assertEqual(download.call_count, 2)
        restore.assert_called_once_with(page, date(2026, 10, 9), date(2026, 10, 19))
        capture.assert_called_once()

    def test_second_download_failure_is_reported_with_current_url(self):
        order = {"order": "W000317", "date": "10/12/2026", "internal_id": "317"}
        page = Mock(url="https://app.kometsales.com/orderSummary/other.do#st")
        with (
            patch.object(stage, "find_filtered_order", side_effect=[order, order]),
            patch.object(stage, "restore_filtered_orders"),
            patch.object(stage, "capture_download_error"),
            patch.object(stage.komet, "download_export_details", side_effect=RuntimeError("no download")) as download,
        ):
            with self.assertRaisesRegex(RuntimeError, "falló dos veces; URL actual"):
                stage.download_with_recovery(page, order, 1, date(2026, 10, 9), date(2026, 10, 19))
        self.assertEqual(download.call_count, 2)

    def test_failed_first_download_does_not_skip_next_order(self):
        orders = [
            {"order": "W000307", "date": "10/09/2026", "internal_id": "307"},
            {"order": "W000317", "date": "10/12/2026", "internal_id": "317"},
        ]
        with tempfile.TemporaryDirectory() as directory:
            output_dir = Path(directory)
            with (
                patch.object(stage, "ARTIFACTS_DIR", output_dir),
                patch.object(stage, "LEDGER_PATH", output_dir / "ledger.json"),
                patch.object(stage.komet, "collect_order_rows", return_value=orders),
                patch.object(stage.komet, "click_next_page", return_value=False),
                patch.object(stage, "download_with_recovery", side_effect=[RuntimeError("no download"), Path("order.xls")]),
                patch.object(stage.transform, "parse_order_xls", return_value=(self.details, [self.box])),
                patch.object(stage.transform, "write_order_workbook", side_effect=lambda template, rows, path: path.write_bytes(b"xlsx")),
            ):
                report = {"orders": []}
                ledger = set()
                stage.transform_visible_orders(
                    None, ledger, set(), self.homologation, self.template, report,
                    date(2026, 10, 9), date(2026, 10, 19),
                )
            self.assertEqual([item["status"] for item in report["orders"]], ["error", "generated"])
            self.assertEqual(ledger, {"W000317|10/12/2026"})

    def test_stage_continues_after_one_bad_order_and_skips_success_on_retry(self):
        orders = [
            {"order": "W000307", "date": "10/09/2026", "internal_id": "307"},
            {"order": "W000317", "date": "10/12/2026", "internal_id": "317"},
        ]
        with tempfile.TemporaryDirectory() as directory:
            output_dir = Path(directory)
            ledger_path = output_dir / "ledger.json"

            def downloaded(_page, order, _index):
                return Path(order["order"] + ".xls")

            def parsed(path):
                number = path.stem[-6:]
                return {**self.details, "order_number": number}, [self.box]

            def rows(details, boxes, lookup):
                if details["order_number"] == "000307":
                    raise TransformationError("Sin homologación")
                return build_order_rows(details, boxes, lookup)

            def written(_template, values, destination):
                destination.write_bytes(b"xlsx-for-test")

            patches = [
                patch.object(stage, "ARTIFACTS_DIR", output_dir),
                patch.object(stage, "LEDGER_PATH", ledger_path),
                patch.object(stage.komet, "collect_order_rows", return_value=orders),
                patch.object(stage, "download_with_recovery", side_effect=lambda page, order, index, start, end: downloaded(page, order, index)),
                patch.object(stage.komet, "click_next_page", return_value=False),
                patch.object(stage.transform, "parse_order_xls", side_effect=parsed),
                patch.object(stage.transform, "build_order_rows", side_effect=rows),
                patch.object(stage.transform, "write_order_workbook", side_effect=written),
            ]
            for current in patches:
                current.start()
            try:
                ledger = set()
                report = {"orders": []}
                stage.transform_visible_orders(None, ledger, set(), self.homologation, self.template, report, date(2026, 10, 9), date(2026, 10, 19))
                self.assertEqual([item["status"] for item in report["orders"]], ["error", "generated"])
                self.assertEqual(load_transformed_keys(ledger_path), {"W000317|10/12/2026"})
                self.assertEqual(len(list(output_dir.glob("*.xlsx"))), 1)

                next_report = {"orders": []}
                stage.transform_visible_orders(None, ledger, set(), self.homologation, self.template, next_report, date(2026, 10, 9), date(2026, 10, 19))
                self.assertEqual([item["status"] for item in next_report["orders"]], ["error", "already_transformed"])
            finally:
                for current in reversed(patches):
                    current.stop()


if __name__ == "__main__":
    unittest.main()
