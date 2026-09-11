from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
import re
from typing import Any, Sequence

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import from_excel


DATE_NUMBER_FORMAT = r"yyyy\-mm\-dd"
KOMET_SHEET_NAME = "Availability"


@dataclass(frozen=True)
class OutputRow:
    source_index: int
    values: tuple[Any, ...]
    is_added: bool = False
    is_blank: bool = False


@dataclass(frozen=True)
class InventoryTransformResult:
    output_rows: tuple[OutputRow, ...]
    original_rows: int
    removed_rows: int
    added_rows: int
    product_count: int
    final_sunday_rows: int
    final_immediate_rows: int
    copied_data_correct: bool


@dataclass(frozen=True)
class InventoryRefreshResult:
    updated_rows: int
    decreased_rows: int
    before_total: float
    after_total: float
    inventory_rows: int
    availability_rows: int
    formula_cells_replaced: int


def next_business_day(value: date) -> date:
    candidate = value + timedelta(days=1)
    while candidate.weekday() == 6:
        candidate += timedelta(days=1)
    return candidate


def business_window_end(assumed_today: date, business_days_after: int = 2) -> date:
    end = assumed_today
    for _ in range(business_days_after):
        end = next_business_day(end)
    return end


def _as_date(value: Any, epoch: Any = None) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        converted = from_excel(value, epoch=epoch) if epoch is not None else from_excel(value)
        return converted.date() if isinstance(converted, datetime) else converted
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
        except ValueError:
            pass
        for pattern in ("%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                continue
    return None


def _has_value(row: Sequence[Any]) -> bool:
    return any(value is not None and value != "" for value in row)


def _komet_sheet(workbook):
    try:
        return workbook[KOMET_SHEET_NAME]
    except KeyError as exc:
        raise RuntimeError(
            f"El XLS de SharePoint no contiene la pestaña requerida {KOMET_SHEET_NAME}."
        ) from exc


def normalize_date_formats(source_path: Path, destination: Path) -> int:
    workbook = load_workbook(source_path, data_only=False)
    date_columns: list[tuple[str, int, int]] = []
    changed = 0
    try:
        worksheet = _komet_sheet(workbook)
        date_column = None
        header_row = None
        for row in worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row)):
            for cell in row:
                if str(cell.value or "").strip().lower() == "available from":
                    date_column = cell.column
                    header_row = cell.row
                    break
            if date_column is not None:
                break
        if date_column is not None and header_row is not None:
            date_columns.append((worksheet.title, header_row, date_column))
            for row in range(header_row + 1, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=date_column)
                if cell.value is not None:
                    cell.number_format = DATE_NUMBER_FORMAT
                    changed += 1
        if changed == 0:
            raise RuntimeError("No se encontraron fechas en la columna Available From del XLS de SharePoint.")
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
    finally:
        workbook.close()

    verification = load_workbook(destination, read_only=True, data_only=False)
    try:
        formats = set()
        for sheet_name, header_row, column in date_columns:
            worksheet = verification[sheet_name]
            for row in range(header_row + 1, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=column)
                if cell.value is not None:
                    formats.add(cell.number_format)
    finally:
        verification.close()
    if formats != {DATE_NUMBER_FORMAT}:
        raise RuntimeError(f"El XLS normalizado conserva formatos de fecha inesperados: {sorted(formats)}")
    return changed


def _same_value(left: Any, right: Any, *, date_value: bool = False) -> bool:
    if date_value:
        return _as_date(left) == _as_date(right)
    if left is None and right == "":
        return True
    if right is None and left == "":
        return True
    return left == right


def apply_inventory_rules(
    rows: Sequence[Sequence[Any]],
    *,
    product_column: int,
    date_column: int,
    assumed_today: date,
    epoch: Any = None,
) -> InventoryTransformResult:
    """Apply the inventory-box date rules to row values.

    Column indexes are zero-based. Blank rows are kept at the end so the
    workbook's existing formatted blank rows remain available.
    """
    window_end = business_window_end(assumed_today)
    original_rows = len(rows)
    retained: list[OutputRow] = []
    blank_rows: list[OutputRow] = []
    product_order: list[str] = []
    seen_products: set[str] = set()
    latest_by_product: dict[str, tuple[date, int, tuple[Any, ...]]] = {}
    removed_rows = 0

    for source_index, raw_row in enumerate(rows):
        values = tuple(raw_row)
        if not _has_value(values):
            blank_rows.append(OutputRow(source_index, values, is_blank=True))
            continue

        product = ""
        if product_column < len(values) and values[product_column] is not None:
            product = str(values[product_column]).strip()
        if product and product not in seen_products:
            seen_products.add(product)
            product_order.append(product)
        available = (
            _as_date(values[date_column], epoch=epoch)
            if date_column < len(values)
            else None
        )

        if available is not None and (
            available.weekday() == 6
            or assumed_today <= available <= window_end
        ):
            removed_rows += 1
            continue

        retained.append(OutputRow(source_index, values))
        if not product:
            continue
        if available is None:
            continue
        previous = latest_by_product.get(product)
        if previous is None or available >= previous[0]:
            latest_by_product[product] = (available, source_index, values)

    missing_products = [product for product in product_order if product not in latest_by_product]
    if missing_products:
        joined = ", ".join(missing_products)
        raise RuntimeError(
            "No quedó una fila con fecha para estas descripciones después de aplicar "
            f"las eliminaciones: {joined}."
        )

    additions: list[OutputRow] = []
    copied_data_correct = True
    for product in product_order:
        latest_date, source_index, source_values = latest_by_product[product]
        new_values = list(source_values)
        new_date = next_business_day(latest_date)
        new_values[date_column] = new_date
        additions.append(
            OutputRow(source_index, tuple(new_values), is_added=True)
        )
        copied_data_correct = copied_data_correct and all(
            _same_value(source_values[index], new_values[index], date_value=index == date_column)
            for index in range(len(source_values))
            if index != date_column
        )

    output_rows = tuple(retained + additions + blank_rows)
    final_dates = [
        _as_date(row.values[date_column], epoch=epoch)
        for row in output_rows
        if not row.is_blank and date_column < len(row.values)
    ]
    final_sunday_rows = sum(
        value is not None and value.weekday() == 6 for value in final_dates
    )
    final_immediate_rows = sum(
        value is not None and assumed_today <= value <= window_end
        for value in final_dates
    )

    return InventoryTransformResult(
        output_rows=output_rows,
        original_rows=original_rows,
        removed_rows=removed_rows,
        added_rows=len(additions),
        product_count=len(product_order),
        final_sunday_rows=final_sunday_rows,
        final_immediate_rows=final_immediate_rows,
        copied_data_correct=copied_data_correct,
    )


def _header_columns(sheet) -> tuple[int, int, int]:
    for row in sheet.iter_rows(min_row=1, max_row=min(20, sheet.max_row)):
        headers = {
            str(cell.value or "").strip().lower(): cell.column - 1
            for cell in row
            if cell.value is not None
        }
        if "product description" in headers and "available from" in headers:
            return row[0].row, headers["product description"], headers["available from"]
    raise RuntimeError(
        "La hoja activa no contiene los encabezados Product Description y Available From."
    )


def _copy_row_format(sheet, source_row: int, target_row: int, max_column: int) -> None:
    if source_row != target_row:
        source_dimension = sheet.row_dimensions[source_row]
        target_dimension = sheet.row_dimensions[target_row]
        target_dimension.height = source_dimension.height
        target_dimension.hidden = source_dimension.hidden
        target_dimension.outlineLevel = source_dimension.outlineLevel
        target_dimension.collapsed = source_dimension.collapsed
    for column in range(1, max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.comment is not None:
            target.comment = copy(source.comment)


def _capture_row_formats(sheet, max_row: int, max_column: int) -> dict[int, dict[str, Any]]:
    formats: dict[int, dict[str, Any]] = {}
    for row in range(1, max_row + 1):
        dimension = sheet.row_dimensions[row]
        cells = []
        for column in range(1, max_column + 1):
            cell = sheet.cell(row, column)
            cells.append(
                (
                    copy(cell._style) if cell.has_style else None,
                    cell.number_format,
                    copy(cell.comment) if cell.comment is not None else None,
                )
            )
        formats[row] = {
            "height": dimension.height,
            "hidden": dimension.hidden,
            "outlineLevel": dimension.outlineLevel,
            "collapsed": dimension.collapsed,
            "cells": tuple(cells),
        }
    return formats


def _copy_captured_row_format(
    sheet,
    captured_formats: dict[int, dict[str, Any]],
    source_row: int,
    target_row: int,
    max_column: int,
) -> None:
    source_format = captured_formats[source_row]
    if source_row != target_row:
        target_dimension = sheet.row_dimensions[target_row]
        target_dimension.height = source_format["height"]
        target_dimension.hidden = source_format["hidden"]
        target_dimension.outlineLevel = source_format["outlineLevel"]
        target_dimension.collapsed = source_format["collapsed"]
    for column, (style, number_format, comment) in enumerate(
        source_format["cells"],
        start=1,
    ):
        target = sheet.cell(target_row, column)
        if style is not None:
            target._style = copy(style)
        if number_format:
            target.number_format = number_format
        if comment is not None:
            target.comment = copy(comment)


def _verify_saved_workbook(
    output_path: Path,
    expected: InventoryTransformResult,
    *,
    header_row: int,
    date_column: int,
    assumed_today: date,
) -> None:
    workbook = load_workbook(output_path, data_only=False)
    try:
        sheet = _komet_sheet(workbook)
        actual_rows = [
            tuple(cell.value for cell in row)
            for row in sheet.iter_rows(
                min_row=header_row + 1,
                max_row=header_row + len(expected.output_rows),
            )
        ]
        if len(actual_rows) != len(expected.output_rows):
            raise RuntimeError("La cantidad de filas del XLS procesado no coincide con la esperada.")
        for row_index, (actual, expected_row) in enumerate(
            zip(actual_rows, expected.output_rows), start=header_row + 1
        ):
            if len(actual) != len(expected_row.values):
                raise RuntimeError(f"La fila {row_index} tiene una cantidad inesperada de columnas.")
            for column, (actual_value, expected_value) in enumerate(
                zip(actual, expected_row.values)
            ):
                if not _same_value(
                    actual_value,
                    expected_value,
                    date_value=column == date_column,
                ):
                    raise RuntimeError(
                        f"El valor de la fila {row_index}, columna {column + 1} "
                        "no coincide con el resultado validado."
                    )

        output_dates = [
            _as_date(row[date_column])
            for row in actual_rows
            if _has_value(row) and date_column < len(row)
        ]
        if any(value is not None and value.weekday() == 6 for value in output_dates):
            raise RuntimeError("El XLS procesado todavía contiene fechas en domingo.")
        window_end = business_window_end(assumed_today)
        if any(
            value is not None and assumed_today <= value <= window_end
            for value in output_dates
        ):
            raise RuntimeError("El XLS procesado todavía contiene fechas dentro de la ventana eliminada.")

        source_target_rows = {
            expected_row.source_index: row_index
            for row_index, expected_row in enumerate(
                expected.output_rows,
                start=header_row + 1,
            )
            if not expected_row.is_added and not expected_row.is_blank
        }
        for row_index, expected_row in enumerate(
            expected.output_rows,
            start=header_row + 1,
        ):
            if not expected_row.is_added:
                continue
            source_row = source_target_rows.get(expected_row.source_index)
            if source_row is None:
                raise RuntimeError(
                    f"No se encontró la fila base retenida para la fila nueva {row_index}."
                )
            for column in range(1, sheet.max_column + 1):
                if sheet.cell(row_index, column)._style != sheet.cell(source_row, column)._style:
                    raise RuntimeError(
                        f"El formato de la fila nueva {row_index} no coincide con su fila base."
                    )
    finally:
        workbook.close()


def _header_map(sheet, required_headers: set[str], *, max_rows: int = 30) -> tuple[int, dict[str, int]]:
    for row in sheet.iter_rows(min_row=1, max_row=min(max_rows, sheet.max_row)):
        headers = {
            str(cell.value or "").strip().lower(): cell.column
            for cell in row
            if cell.value is not None and str(cell.value).strip()
        }
        if required_headers.issubset(headers):
            return row[0].row, headers
    raise RuntimeError(
        "No se encontraron las columnas requeridas: "
        + ", ".join(sorted(required_headers))
        + "."
    )


def _table_bounds(sheet, table_name: str) -> tuple[int, int, int, int]:
    try:
        table = sheet.tables[table_name]
    except KeyError as exc:
        raise RuntimeError(f"La hoja {sheet.title} no contiene la tabla {table_name}.") from exc
    return range_boundaries(table.ref)


def _normalized_product(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()[:30]


def _numeric_quantity(value: Any) -> float | None:
    if isinstance(value, bool) or value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def _inventory_rows(workbook, sheet_name: str, *, table_name: str | None = None) -> tuple[list[dict[str, Any]], int, dict[str, int]]:
    sheet = workbook[sheet_name]
    if table_name:
        min_col, min_row, max_col, max_row = _table_bounds(sheet, table_name)
        header_row = min_row
        headers = {
            str(sheet.cell(header_row, column).value or "").strip().lower(): column
            for column in range(min_col, max_col + 1)
        }
    else:
        header_row, headers = _header_map(sheet, {"product", "qty", "aging"})
        min_col = min(headers.values())
        max_col = max(headers.values())
        max_row = sheet.max_row

    rows: list[dict[str, Any]] = []
    for row_number in range(header_row + 1, max_row + 1):
        values = {
            name: sheet.cell(row_number, column).value
            for name, column in headers.items()
            if min_col <= column <= max_col
        }
        if not any(value is not None and value != "" for value in values.values()):
            continue
        if not str(values.get("product") or "").strip():
            continue
        rows.append(values)
    return rows, header_row, headers


def _inventory_rows_from_xls(path: Path) -> tuple[list[dict[str, Any]], int, dict[str, int]]:
    """Read Komet's legacy binary Excel export without treating it as XLSX."""
    frame = pd.read_excel(path, header=None, engine="xlrd")
    values = frame.astype(object).where(pd.notna(frame), None).values.tolist()
    required_headers = {"product", "qty", "aging"}
    for row_index, row_values in enumerate(values[:30], start=1):
        headers = {
            str(value).strip().lower(): column_index
            for column_index, value in enumerate(row_values, start=1)
            if value is not None and str(value).strip()
        }
        if not required_headers.issubset(headers):
            continue
        rows: list[dict[str, Any]] = []
        for row_values in values[row_index:]:
            row = {name: row_values[column - 1] if column <= len(row_values) else None for name, column in headers.items()}
            if not any(value is not None and value != "" for value in row.values()):
                continue
            if not str(row.get("product") or "").strip():
                continue
            rows.append(row)
        return rows, row_index, headers
    raise RuntimeError("El XLS de Komet no contiene las columnas Product, Qty y Aging.")


def _inventory_key(row: dict[str, Any], assumed_today: date, *, use_awb_date: bool = False) -> tuple[str, date] | None:
    product = _normalized_product(row.get("product"))
    aging = _numeric_quantity(row.get("aging"))
    if not product or aging is None:
        return None
    base_date = assumed_today
    if use_awb_date:
        awb = str(row.get("awb") or "")
        match = re.search(r"((?:19|20)\d{2})[-/]?(\d{2})[-/]?(\d{2})", awb)
        if match:
            base_date = date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    return product, base_date + timedelta(days=abs(int(aging)))


def _aggregate_inventory(rows: list[dict[str, Any]], assumed_today: date, *, use_awb_date: bool = False) -> dict[tuple[str, date], float]:
    totals: dict[tuple[str, date], float] = {}
    for row in rows:
        key = _inventory_key(row, assumed_today, use_awb_date=use_awb_date)
        quantity = _numeric_quantity(row.get("qty"))
        if key is None or quantity is None:
            continue
        totals[key] = totals.get(key, 0) + quantity
    return totals


def _availability_current_quantity(
    value: Any,
    row: dict[str, Any],
    old_inventory_totals: dict[tuple[str, date], float],
    assumed_today: date,
) -> float:
    quantity = _numeric_quantity(value)
    if quantity is not None:
        return quantity
    key = (row["product_key"], row["available"])
    if key in old_inventory_totals:
        return old_inventory_totals[key]
    raise RuntimeError(
        "No se pudo obtener la cantidad actual para una fórmula de Availability: "
        f"producto={row['product']!r}, fecha={row['available']}.")


def refresh_workbook_with_komet_inventory(
    source_path: Path,
    komet_inventory_path: Path,
    output_path: Path,
    *,
    assumed_today: date,
) -> InventoryRefreshResult:
    workbook = load_workbook(source_path, data_only=False, keep_links=True)
    try:
        availability = workbook[KOMET_SHEET_NAME]
        _, availability_header_row, availability_max_col, availability_max_row = _table_bounds(
            availability, "tblAvailability"
        )
        availability_headers = {
            str(availability.cell(availability_header_row, column).value or "").strip().lower(): column
            for column in range(1, availability_max_col + 1)
        }
        required = {"product description", "qty packages", "available from"}
        if not required.issubset(availability_headers):
            raise RuntimeError("tblAvailability no contiene Product Description, Qty Packages y Available From.")

        old_inventory_book = workbook
        old_inventory_rows, _, _ = _inventory_rows(old_inventory_book, "Inventory", table_name="tblInventory")
        if komet_inventory_path.suffix.lower() == ".xls":
            new_inventory_rows, _, _ = _inventory_rows_from_xls(komet_inventory_path)
        else:
            komet_book = load_workbook(komet_inventory_path, data_only=False, keep_links=True)
            try:
                source_sheet_name = komet_book.sheetnames[0]
                new_inventory_rows, _, _ = _inventory_rows(komet_book, source_sheet_name)
            finally:
                komet_book.close()

        old_inventory_totals = _aggregate_inventory(old_inventory_rows, assumed_today, use_awb_date=True)
        new_inventory_totals = _aggregate_inventory(new_inventory_rows, assumed_today)
        _, new_inventory_rows = _write_inventory_rows_from_rows(
            workbook,
            new_inventory_rows,
        )

        before_total = 0.0
        after_total = 0.0
        updated_rows = 0
        decreased_rows = 0
        formula_cells_replaced = 0
        for row_number in range(availability_header_row + 1, availability_max_row + 1):
            product = availability.cell(row_number, availability_headers["product description"]).value
            available_value = availability.cell(row_number, availability_headers["available from"]).value
            available = _as_date(available_value, epoch=workbook.epoch)
            if not product or available is None:
                continue
            product_key = _normalized_product(product)
            current_cell = availability.cell(row_number, availability_headers["qty packages"])
            current = _availability_current_quantity(
                current_cell.value,
                {"product": product, "product_key": product_key, "available": available},
                old_inventory_totals,
                assumed_today,
            )
            inventory_quantity = new_inventory_totals.get((product_key, available), 0.0)
            final_quantity = min(current, inventory_quantity)
            before_total += current
            after_total += final_quantity
            if final_quantity != current:
                updated_rows += 1
            if final_quantity < current:
                decreased_rows += 1
            if current_cell.value.__class__.__name__ == "ArrayFormula" or (
                isinstance(current_cell.value, str) and current_cell.value.startswith("=")
            ):
                formula_cells_replaced += 1
            current_cell.value = int(final_quantity) if final_quantity.is_integer() else final_quantity

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
    finally:
        workbook.close()

    verification = load_workbook(output_path, data_only=False)
    try:
        availability = verification[KOMET_SHEET_NAME]
        inventory = verification["Inventory"]
        _, _, _, inventory_end = _table_bounds(inventory, "tblInventory")
        _, inventory_header_row, _, _ = _table_bounds(inventory, "tblInventory")
        if inventory_end != inventory_header_row + len(new_inventory_rows):
            raise RuntimeError("tblInventory no fue ampliada al número exacto de filas descargadas.")
        qty_column = availability_headers["qty packages"]
        for row_number in range(availability_header_row + 1, availability_max_row + 1):
            value = availability.cell(row_number, qty_column).value
            if (
                isinstance(value, str) and value.startswith("=")
            ) or value.__class__.__name__ == "ArrayFormula":
                raise RuntimeError(f"Availability!{get_column_letter(qty_column)}{row_number} conserva una fórmula.")
    finally:
        verification.close()

    return InventoryRefreshResult(
        updated_rows=updated_rows,
        decreased_rows=decreased_rows,
        before_total=before_total,
        after_total=after_total,
        inventory_rows=len(new_inventory_rows),
        availability_rows=availability_max_row - availability_header_row,
        formula_cells_replaced=formula_cells_replaced,
    )


def _write_inventory_rows_from_rows(workbook, source_rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    destination = workbook["Inventory"]
    min_col, min_row, max_col, old_max_row = _table_bounds(destination, "tblInventory")
    destination_headers = {
        str(destination.cell(min_row, column).value or "").strip().lower(): column
        for column in range(min_col, max_col + 1)
    }
    old_rows = [
        {name: destination.cell(row, column).value for name, column in destination_headers.items()}
        for row in range(min_row + 1, old_max_row + 1)
    ]
    clear_until = max(old_max_row, min_row + len(source_rows))
    for row in range(min_row + 1, clear_until + 1):
        for column in range(min_col, max_col + 1):
            destination.cell(row, column).value = None
    template_row = min_row + 1 if old_max_row >= min_row + 1 else min_row
    for offset, source_row in enumerate(source_rows, start=1):
        target_row = min_row + offset
        if target_row != template_row:
            _copy_row_format(destination, template_row, target_row, max_col)
        for name, column in destination_headers.items():
            destination.cell(target_row, column).value = source_row.get(name)
    table = destination.tables["tblInventory"]
    table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{min_row + len(source_rows)}"
    if table.autoFilter:
        table.autoFilter.ref = table.ref
    return old_rows, source_rows


def transform_inventory_workbook(
    source_path: Path,
    output_path: Path,
    *,
    assumed_today: date,
) -> InventoryTransformResult:
    workbook = load_workbook(source_path, data_only=False, keep_links=True)
    try:
        sheet = _komet_sheet(workbook)
        header_row, product_column, date_column = _header_columns(sheet)
        data_start = header_row + 1
        original_max_row = sheet.max_row
        max_column = sheet.max_column
        captured_formats = _capture_row_formats(sheet, original_max_row, max_column)
        rows = [
            tuple(cell.value for cell in sheet.iter_rows(min_row=row, max_row=row, max_col=max_column).__next__())
            for row in range(data_start, original_max_row + 1)
        ]
        result = apply_inventory_rules(
            rows,
            product_column=product_column,
            date_column=date_column,
            assumed_today=assumed_today,
            epoch=workbook.epoch,
        )

        for target_offset, output_row in enumerate(result.output_rows):
            target_row = data_start + target_offset
            source_row = data_start + output_row.source_index
            _copy_captured_row_format(
                sheet,
                captured_formats,
                source_row,
                target_row,
                max_column,
            )
            for column, value in enumerate(output_row.values, start=1):
                sheet.cell(target_row, column).value = value

        final_row = data_start + len(result.output_rows) - 1
        for row in range(final_row + 1, original_max_row + 1):
            for column in range(1, max_column + 1):
                sheet.cell(row, column).value = None

        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
    finally:
        workbook.close()

    _verify_saved_workbook(
        output_path,
        result,
        header_row=header_row,
        date_column=date_column,
        assumed_today=assumed_today,
    )
    return result


def create_single_sheet_workbook(source_path: Path, output_path: Path) -> None:
    """Save a copy containing only the Availability sheet for Kometsales."""
    workbook = load_workbook(source_path, data_only=False, keep_links=True)
    try:
        target_sheet = _komet_sheet(workbook)
        for worksheet in tuple(workbook.worksheets):
            if worksheet is not target_sheet:
                workbook.remove(worksheet)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
    finally:
        workbook.close()

    verification = load_workbook(output_path, data_only=False, read_only=True)
    try:
        if verification.sheetnames != [target_sheet.title]:
            raise RuntimeError("La copia para Kometsales no conserva únicamente la pestaña Availability.")
    finally:
        verification.close()
