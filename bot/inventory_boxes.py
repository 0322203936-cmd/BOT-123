from __future__ import annotations

import os
import re
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import Page, TimeoutError as PlaywrightTimeoutError, sync_playwright

from sharepoint_sync import (
    download_sharepoint_file,
    graph_token,
    resolve_sharepoint_item_by_url,
)


KOMET_LOGIN_URL = "https://app.kometsales.com/sign-in/login.do#st"
KOMET_BOXES_URL = "https://app.kometsales.com/inventory-pricing/list_pricing.do#st"
SHAREPOINT_BOXES_URL = (
    "https://pacificafarms.sharepoint.com/:x:/r/sites/"
    "requerimientovsproyeccion/_layouts/15/Doc.aspx?"
    "sourcedoc=%7BC0D676CF-1FBB-4922-88D2-FE4D6FD4526A%7D&"
    "file=Inventory%20Upload%20Boxes%20050926.xlsx&action=default&mobileredirect=true"
)

ARTIFACTS_DIR = Path("artifacts/inventory_boxes")
CAPTURES_DIR = ARTIFACTS_DIR / "capturas"
REPORTS_DIR = ARTIFACTS_DIR / "reportes"
SOURCE_FILENAME = "inventory-upload-boxes-source.xlsx"
DATE_NUMBER_FORMAT = r"yyyy\-mm\-dd"
MAX_DELETE_BATCHES = 50


def required_secret(name: str) -> str:
    value = os.environ.get(name, "").strip()
    if not value:
        raise RuntimeError(f"Falta configurar el secreto {name}.")
    return value


def capture(page: Page, filename: str) -> None:
    CAPTURES_DIR.mkdir(parents=True, exist_ok=True)
    page.screenshot(path=str(CAPTURES_DIR / filename), full_page=True)
    print(f"Captura guardada: {filename}", flush=True)


def visible_locator(locator) -> object | None:
    for index in range(locator.count()):
        candidate = locator.nth(index)
        try:
            if candidate.is_visible():
                return candidate
        except Exception:
            continue
    return None


def click_first_visible(page: Page, locators: list, description: str) -> None:
    for locator in locators:
        candidate = visible_locator(locator)
        if candidate is None:
            continue
        try:
            candidate.click(timeout=15_000)
            print(f"Clic: {description}", flush=True)
            return
        except PlaywrightTimeoutError:
            continue
    raise RuntimeError(f"No se encontró el control visible: {description}.")


def click_text(page: Page, text: str, description: str | None = None) -> None:
    pattern = re.compile(rf"^\s*{re.escape(text)}\s*$", re.IGNORECASE)
    click_first_visible(
        page,
        [
            page.get_by_role("button", name=pattern),
            page.get_by_role("link", name=pattern),
            page.get_by_text(pattern),
        ],
        description or text,
    )


def wait_for_network(page: Page, timeout: int = 60_000) -> None:
    try:
        page.wait_for_load_state("networkidle", timeout=timeout)
    except PlaywrightTimeoutError:
        print(f"Aviso: la página no llegó a networkidle. URL actual: {page.url}", flush=True)


def login_kometsales(page: Page, user: str, password: str) -> None:
    print("Abriendo Kometsales...", flush=True)
    page.goto(KOMET_LOGIN_URL, wait_until="domcontentloaded", timeout=60_000)
    page.locator(
        'input[type="email"], input[placeholder*="usuario" i], input[name*="user" i], input[type="text"]'
    ).first.fill(user)
    page.locator(
        'input[type="password"], input[placeholder*="contraseña" i], input[placeholder*="password" i]'
    ).first.fill(password)
    click_first_visible(
        page,
        [
            page.get_by_role("button", name=re.compile(r"iniciar sesión|entrar|login", re.I)),
            page.locator('input[type="submit"]'),
        ],
        "Iniciar sesión",
    )
    wait_for_network(page)
    page.wait_for_timeout(2_000)

    account = page.get_by_text(re.compile(r"^\s*PACIFICA\s+FARMS\s*-\s*CAL\s*$", re.I))
    if account.count() > 0 and visible_locator(account) is not None:
        click_first_visible(page, [account], "PACIFICA FARMS - CAL")
        wait_for_network(page)
        page.wait_for_timeout(2_000)

    if "accounts.do" in page.url.lower():
        raise RuntimeError("Kometsales mostró la selección de cuenta, pero no se pudo elegir PACIFICA FARMS - CAL.")
    print(f"Sesión de Kometsales iniciada. URL: {page.url}", flush=True)


def open_boxes(page: Page) -> None:
    print("Abriendo Inventario > Cajas...", flush=True)
    try:
        click_first_visible(
            page,
            [
                page.get_by_role("button", name=re.compile(r"inventario", re.I)),
                page.get_by_text(re.compile(r"^\s*INVENTARIO\s*$", re.I)),
            ],
            "Inventario",
        )
        page.wait_for_timeout(700)
        click_text(page, "Cajas")
        page.wait_for_url("**/inventory-pricing/list_pricing.do**", timeout=30_000)
    except (PlaywrightTimeoutError, RuntimeError):
        print("Aviso: no se pudo navegar por el menú; abriendo Cajas por su ruta directa.", flush=True)
        page.goto(KOMET_BOXES_URL, wait_until="domcontentloaded", timeout=60_000)
    wait_for_network(page)
    page.wait_for_timeout(1_500)
    capture(page, "01_cajas.png")


def inventory_is_empty(page: Page) -> bool:
    no_records = page.locator("#tdLabelNoRecords")
    if visible_locator(no_records) is not None:
        return True
    empty_message = page.get_by_text(re.compile(r"Parece que no podemos encontrar ningún registro", re.I))
    return visible_locator(empty_message) is not None


def awb_checkbox(page: Page):
    header_pattern = re.compile(r"^\s*AWB\s*$", re.I)
    candidates = [
        page.locator("th").filter(has_text=header_pattern).locator('input[type="checkbox"]'),
        page.locator('[role="columnheader"]').filter(has_text=header_pattern).locator('input[type="checkbox"]'),
        page.locator("label").filter(has_text=header_pattern).locator('input[type="checkbox"]'),
        page.locator('input[type="checkbox"][aria-label*="AWB" i]'),
        page.locator('input[type="checkbox"][title*="AWB" i]'),
    ]
    for candidate in candidates:
        found = visible_locator(candidate)
        if found is not None:
            return found

    awb_text = page.get_by_text(header_pattern)
    for index in range(awb_text.count()):
        label = awb_text.nth(index)
        for ancestor in [
            label.locator("xpath=ancestor::th[1]"),
            label.locator("xpath=ancestor::*[@role='columnheader'][1]"),
            label.locator("xpath=ancestor::tr[1]"),
            label.locator("xpath=.."),
        ]:
            found = visible_locator(ancestor.locator('input[type="checkbox"]'))
            if found is not None:
                return found
    return None


def select_all_inventory(page: Page) -> None:
    checkbox = awb_checkbox(page)
    if checkbox is None:
        raise RuntimeError("No se encontró la casilla para seleccionar todas las cajas junto a AWB.")
    if not checkbox.is_checked():
        checkbox.click(timeout=15_000)
    page.wait_for_timeout(700)
    print("Todas las cajas visibles fueron seleccionadas mediante AWB.", flush=True)


def confirm_mass_delete(page: Page) -> None:
    click_first_visible(
        page,
        [
            page.locator("#inventoryPricingActions"),
            page.get_by_role("button", name=re.compile(r"acciones", re.I)),
        ],
        "Acciones",
    )
    click_first_visible(
        page,
        [
            page.get_by_text(re.compile(r"Borrar\s+inventario\s+masivamente", re.I)),
            page.get_by_role("menuitem", name=re.compile(r"Borrar\s+inventario\s+masivamente", re.I)),
        ],
        "Borrar inventario masivamente",
    )
    page.wait_for_timeout(500)

    dialog = page.get_by_role("dialog")
    input_candidates = [
        dialog.locator('input[type="text"], input:not([type])'),
        page.locator('input[placeholder*="CONFIRMAR" i], input[name*="confirm" i]'),
    ]
    confirmation_prompt = page.get_by_text(re.compile(r"CONFIRMAR", re.I))
    for index in range(confirmation_prompt.count()):
        prompt = confirmation_prompt.nth(index)
        input_candidates.extend(
            [
                prompt.locator("xpath=ancestor::*[@role='dialog'][1]").locator(
                    'input[type="text"], input:not([type])'
                ),
                prompt.locator("xpath=..").locator('input[type="text"], input:not([type])'),
                prompt.locator("xpath=../..").locator('input[type="text"], input:not([type])'),
            ]
        )
    confirmation_input = None
    for locator in input_candidates:
        confirmation_input = visible_locator(locator)
        if confirmation_input is not None:
            break
    if confirmation_input is None:
        raise RuntimeError("La ventana de borrado masivo no mostró el campo de confirmación.")
    confirmation_input.fill("CONFIRMAR")
    click_first_visible(
        page,
        [
            dialog.get_by_role("button", name=re.compile(r"continuar|confirmar|borrar|eliminar|aceptar", re.I)),
            page.get_by_role("button", name=re.compile(r"continuar|confirmar|borrar|eliminar|aceptar", re.I)),
            page.get_by_role("link", name=re.compile(r"continuar|confirmar|borrar|eliminar|aceptar", re.I)),
            page.locator('input[type="button"], input[type="submit"]'),
        ],
        "Confirmar borrado masivo",
    )


def delete_all_inventory(page: Page) -> None:
    for batch in range(1, MAX_DELETE_BATCHES + 1):
        if inventory_is_empty(page):
            print(f"Inventario vacío después de {batch - 1} lote(s).", flush=True)
            return
        print(f"Borrado masivo lote {batch}...", flush=True)
        select_all_inventory(page)
        confirm_mass_delete(page)
        wait_for_network(page)
        page.wait_for_timeout(3_000)
        capture(page, f"02_borrado_{batch:02d}.png")
    raise RuntimeError(f"El inventario no quedó vacío después de {MAX_DELETE_BATCHES} lotes.")


def upload_boxes(page: Page, workbook_path: Path) -> None:
    print(f"Abriendo Subir XLS cajas: {workbook_path.name}", flush=True)
    click_first_visible(
        page,
        [
            page.get_by_role("link", name=re.compile(r"Subir\s+XLS\s+cajas", re.I)),
            page.get_by_text(re.compile(r"Subir\s+XLS\s+cajas", re.I)),
        ],
        "Subir XLS cajas",
    )
    page.wait_for_url("**/inventory-boxes/upload.do**", timeout=30_000)
    wait_for_network(page)

    file_input = page.locator('input[type="file"]')
    if file_input.count() == 0:
        raise RuntimeError("La página de carga no contiene un selector de archivos.")
    file_input.first.set_input_files(str(workbook_path))
    page.wait_for_timeout(1_000)
    capture(page, "03_archivo_seleccionado.png")

    click_first_visible(
        page,
        [
            page.get_by_role("button", name=re.compile(r"subir|upload|cargar", re.I)),
            page.locator('input[type="submit"]'),
        ],
        "Subir archivo de cajas",
    )
    success = re.compile(r"(subid|cargad|exitos|correctamente|procesad)", re.I)
    try:
        page.get_by_text(success).first.wait_for(state="visible", timeout=20_000)
    except PlaywrightTimeoutError as exc:
        error_message = re.compile(
            r"(error|inv[aá]lid|no se pudo|rechaz|failed|incorrect|must be|required)",
            re.I,
        )
        error_candidates = [
            page.locator(".alert-danger:visible, .error:visible, .fieldError:visible"),
            page.get_by_text(error_message),
        ]
        if any(visible_locator(locator) is not None for locator in error_candidates):
            raise RuntimeError("Kometsales mostró un error después de intentar cargar el XLS.") from exc
        print(
            "Aviso: Kometsales no mostró un mensaje final, pero tampoco mostró un error. "
            "Se da por terminada la carga después de esperar la respuesta de la página.",
            flush=True,
        )
    capture(page, "04_carga_completada.png")
    print("XLS de cajas cargado correctamente en Kometsales.", flush=True)


def download_and_normalize_source() -> Path:
    token = graph_token()
    item = resolve_sharepoint_item_by_url(token, SHAREPOINT_BOXES_URL)
    source_path = download_sharepoint_file(token, item, SOURCE_FILENAME)
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    destination = REPORTS_DIR / f"inventory-upload-boxes-{date.today().isoformat()}.xlsx"

    changed = normalize_date_formats(source_path, destination)
    print(f"Archivo de SharePoint descargado y normalizado: {destination} ({changed} fechas)", flush=True)
    return destination


def normalize_date_formats(source_path: Path, destination: Path) -> int:
    workbook = load_workbook(source_path, data_only=False)
    date_columns: list[tuple[str, int, int]] = []
    changed = 0
    try:
        for worksheet in workbook.worksheets:
            date_column = None
            header_row = None
            for row in worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row)):
                for cell in row:
                    if str(cell.value or "").strip().lower() == "available from":
                        date_column = cell.column
                        header_row = cell.row
                        break
                if date_column is not None:
                    break
            if date_column is None or header_row is None:
                continue
            date_columns.append((worksheet.title, header_row, date_column))
            for row in range(header_row + 1, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=date_column)
                if cell.value is not None:
                    cell.number_format = DATE_NUMBER_FORMAT
                    changed += 1
        if changed == 0:
            raise RuntimeError("No se encontraron fechas en la columna Available From del XLS de SharePoint.")
        workbook.save(destination)
    finally:
        workbook.close()

    verification = load_workbook(destination, read_only=True, data_only=False)
    try:
        formats = set()
        for sheet_name, header_row, column in date_columns:
            worksheet = verification[sheet_name]
            for row in range(header_row + 1, worksheet.max_row + 1):
                cell = worksheet.cell(row=row, column=column)
                if cell.value is not None:
                    formats.add(cell.number_format)
    finally:
        verification.close()
    if formats != {DATE_NUMBER_FORMAT}:
        raise RuntimeError(f"El XLS normalizado conserva formatos de fecha inesperados: {sorted(formats)}")
    return changed


def run() -> None:
    komet_user = required_secret("KOMET_USER")
    komet_password = required_secret("KOMET_PASSWORD")
    source_path = download_and_normalize_source()

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        context = browser.new_context(viewport={"width": 1920, "height": 1080}, accept_downloads=True)
        page = context.new_page()
        try:
            login_kometsales(page, komet_user, komet_password)
            capture(page, "00_sesion_iniciada.png")
            open_boxes(page)
            delete_all_inventory(page)
            upload_boxes(page, source_path)
            print(f"Proceso completo. URL final: {page.url}", flush=True)
        except Exception:
            capture(page, "99_error.png")
            raise
        finally:
            context.close()
            browser.close()


if __name__ == "__main__":
    run()
